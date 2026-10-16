from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

F='Arial'; NAVY='0B4F6C'; TEAL='0B7285'; SOFT='E3F2F5'; INPUT='FFFBE6'; GREY='F2F4F7'
thin=Side(style='thin',color='B8C4CE'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
def font(**k): d=dict(name=F,size=10); d.update(k); return Font(**d)
def fill(c): return PatternFill('solid',fgColor=c)
CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True)
LEFT=Alignment(horizontal='left',vertical='center',wrap_text=False)
N=1000                      # số dòng nhập sẵn định dạng
wb=Workbook()

# ================= Sheet 1: BẢNG CÂN =================
ws=wb.active; ws.title='BANG CAN'
HDR=7; R0=HDR+1; R1=HDR+N
cols=[('STT',6),('Số phiếu',12),('Ngày cân',12),('Giờ cân vào',10),('Giờ cân ra',10),('Biển số xe',13),
      ('Cân lần 1 (xe không) (kg)',14),('Cân lần 2 (có hàng) (kg)',14),('Khối lượng hàng (kg)',14),('Kho xuất',16),('Ghi chú',22),('Kiểm tra DS xe',18)]
for i,(h,w) in enumerate(cols,1): ws.column_dimensions[L(i)].width=w
ws.merge_cells('A1:L1'); ws['A1']='BẢNG CÂN XUẤT HÀNG'; ws['A1'].font=font(size=16,bold=True,color=NAVY); ws['A1'].alignment=CENTER; ws.row_dimensions[1].height=30
ws.merge_cells('A2:L2'); ws['A2']='Mỗi file = 1 đợt xuất hàng. Danh sách xe & tài xế của đợt này nằm ở sheet "DS XE - TAI XE" trong CÙNG file — chỉ áp dụng cho bảng cân này.'
ws['A2'].font=font(size=9,italic=True,color='5B6B7C'); ws['A2'].alignment=Alignment(horizontal='center')
info=[('A3','Đơn vị xuất:','C3'),('A4','Tàu / chuyến:','C4'),('A5','Kho xuất (bàn cân):','C5')]
for a,t,c in info:
    ws[a]=t; ws[a].font=font(bold=True); ws.merge_cells(f'{c}:F{c[1:]}'); ws[c].fill=fill(INPUT); ws[c].border=B; ws[c].font=font(color='0000FF')
    for col in 'CDEF': ws[f'{col}{c[1:]}'].border=B
sumr=[('H3','Số phiếu:',f'=COUNT(I{R0}:I{R1})','#,##0'),('H4','Tổng khối lượng (kg):',f'=SUM(I{R0}:I{R1})','#,##0'),('H5','Tổng khối lượng (tấn):',f'=ROUND(SUM(I{R0}:I{R1})/1000,2)','#,##0.00')]
for a,t,fml,fmt in sumr:
    r=a[1:]; ws.merge_cells(f'H{r}:I{r}'); ws[a]=t; ws[a].font=font(bold=True); ws[a].alignment=Alignment(horizontal='right')
    ws[f'J{r}']=fml; ws[f'J{r}'].number_format=fmt; ws[f'J{r}'].font=font(bold=True,color=TEAL); ws[f'J{r}'].fill=fill(GREY); ws[f'J{r}'].border=B
ws[f'K3']='Từ ngày:'; ws['K3'].font=font(bold=True); ws['K3'].alignment=Alignment(horizontal='right')
ws['L3']=f'=IF(COUNT(C{R0}:C{R1})=0,"",MIN(C{R0}:C{R1}))'; ws['K4']='Đến ngày:'; ws['K4'].font=font(bold=True); ws['K4'].alignment=Alignment(horizontal='right')
ws['L4']=f'=IF(COUNT(C{R0}:C{R1})=0,"",MAX(C{R0}:C{R1}))'
ws['K5']='Chưa có trong DS xe:'; ws['K5'].font=font(bold=True); ws['K5'].alignment=Alignment(horizontal='right')
ws['L5']=f'=COUNTIF(L{R0}:L{R1},"Chưa có*")'
for c in ('L3','L4','L5'): ws[c].font=font(bold=True,color=TEAL); ws[c].fill=fill(GREY); ws[c].border=B
ws['L3'].number_format=ws['L4'].number_format='dd/mm/yyyy'
for i,(h,w) in enumerate(cols,1):
    c=ws.cell(HDR,i,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(TEAL); c.alignment=CENTER; c.border=B
ws.row_dimensions[HDR].height=36
ws.cell(HDR,4).comment=Comment('Gõ giờ dạng 07:05 hoặc 07:05:30 (ô định dạng chữ, giữ nguyên như phiếu cân).','HAK')
ws.cell(HDR,9).comment=Comment('Tự tính = Cân lần 2 − Cân lần 1. Nếu phần mềm cân xuất sẵn khối lượng hàng thì dán đè số vào cột này.','HAK')
ws.cell(HDR,12).comment=Comment('Tự kiểm tra biển số có trong sheet "DS XE - TAI XE" của đợt này không. Cột này chỉ để xem, không cần nhập.','HAK')
DS="'DS XE - TAI XE'"
for r in range(R0,R1+1):
    ws.cell(r,1,f'=IF(F{r}="","",COUNTA($F${R0}:F{r}))')
    ws.cell(r,9,f'=IF(OR(G{r}="",H{r}=""),"",H{r}-G{r})')
    ws.cell(r,12,f'=IF(F{r}="","",IF(COUNTIF({DS}!$B$6:$B$505,F{r})>0,"✓ Có","Chưa có trong DS xe"))')
    for ci in range(1,13):
        c=ws.cell(r,ci); c.border=B; c.font=font(color='0000FF' if ci in (2,3,4,5,6,7,8,10,11) else '000000')
        if ci in (2,3,4,5,6,7,8,10,11): c.fill=fill(INPUT)
    ws.cell(r,3).number_format='dd/mm/yyyy'; ws.cell(r,4).number_format=ws.cell(r,5).number_format='@'; ws.cell(r,4).alignment=ws.cell(r,5).alignment=Alignment(horizontal='center')
    for ci in (7,8,9): ws.cell(r,ci).number_format='#,##0'
    ws.cell(r,1).alignment=Alignment(horizontal='center'); ws.cell(r,6).alignment=Alignment(horizontal='center')
    ws.cell(r,12).font=font(size=9,color='2F9E44')
ws.conditional_formatting.add(f'L{R0}:L{R1}',FormulaRule(formula=[f'LEFT(L{R0},5)="Chưa "'],font=Font(name=F,color='C92A2A',bold=True),fill=fill('FDECEC')))
ws.conditional_formatting.add(f'I{R0}:I{R1}',FormulaRule(formula=[f'AND(I{R0}<>"",I{R0}<=0)'],font=Font(name=F,color='C92A2A',bold=True),fill=fill('FDECEC')))
dv=DataValidation(type='list',formula1=f"={DS}!$B$6:$B$505",allow_blank=True,errorStyle='warning',showErrorMessage=True,
    errorTitle='Biển số chưa có trong DS xe',error='Biển số này chưa có trong sheet "DS XE - TAI XE" của đợt. Bấm Yes để vẫn nhập, sau đó bổ sung vào danh sách xe.',
    promptTitle='Biển số xe',prompt='Chọn từ danh sách xe của đợt, hoặc gõ biển số.',showInputMessage=True)
ws.add_data_validation(dv); dv.add(f'F{R0}:F{R1}')
dvd=DataValidation(type='date',operator='greaterThan',formula1='36526',allow_blank=True,errorStyle='stop',showErrorMessage=True,error='Nhập ngày dạng dd/mm/yyyy.',errorTitle='Ngày cân')
ws.add_data_validation(dvd); dvd.add(f'C{R0}:C{R1}')
dvn=DataValidation(type='decimal',operator='between',formula1='0',formula2='200000',allow_blank=True,errorStyle='stop',showErrorMessage=True,error='Khối lượng cân tính bằng kg (0 – 200.000).',errorTitle='Số cân (kg)')
ws.add_data_validation(dvn); dvn.add(f'G{R0}:H{R1}')
ws.freeze_panes=f'C{R0}'; ws.auto_filter.ref=f'A{HDR}:L{R1}'
ws.print_title_rows=f'{HDR}:{HDR}'; ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=9; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_options.horizontalCentered=True
ws.page_margins.left=ws.page_margins.right=0.4; ws.oddFooter.center.text='Trang &P / &N'; ws.sheet_properties.tabColor=TEAL
ws.sheet_view.zoomScale=100

# ================= Sheet 2: DS XE - TAI XE =================
v=wb.create_sheet('DS XE - TAI XE'); HD=5; V0=6; V1=505
vc=[('STT',6),('Biển số xe',14),('Họ và tên tài xế',26),('Số điện thoại',14),('Mã ĐVVC',11),('Tên đơn vị vận chuyển',44),('Số chuyến trong bảng cân',13),('Ghi chú',20)]
for i,(h,w) in enumerate(vc,1): v.column_dimensions[L(i)].width=w
v.merge_cells('A1:H1'); v['A1']='DANH SÁCH XE VÀ TÀI XẾ VẬN CHUYỂN'; v['A1'].font=font(size=16,bold=True,color=NAVY); v['A1'].alignment=CENTER; v.row_dimensions[1].height=30
v.merge_cells('A2:H2'); v['A2']='Chỉ có hiệu lực cho bảng cân (sheet "BANG CAN") trong cùng file này — không áp dụng cho lô, tàu, đợt khác. Mỗi đợt xuất hàng lập 1 danh sách riêng.'
v['A2'].font=font(size=9,italic=True,color='C92A2A'); v['A2'].alignment=Alignment(horizontal='center',wrap_text=True); v.row_dimensions[2].height=26
v['A3']='Áp dụng cho tàu / đợt:'; v['A3'].font=font(bold=True); v.merge_cells('C3:F3'); v['C3']="=IF('BANG CAN'!C4=\"\",\"(ghi tàu / chuyến ở sheet BANG CAN)\",'BANG CAN'!C4)"
v['C3'].font=font(bold=True,color=TEAL); v['C3'].fill=fill(GREY)
v['G3']='Số lượng xe:'; v['G3'].font=font(bold=True); v['G3'].alignment=Alignment(horizontal='right'); v['H3']=f'=COUNTA(B{V0}:B{V1})'; v['H3'].font=font(bold=True,color=TEAL); v['H3'].fill=fill(GREY)
for i,(h,w) in enumerate(vc,1):
    c=v.cell(HD,i,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(NAVY); c.alignment=CENTER; c.border=B
v.row_dimensions[HD].height=34
v.cell(HD,5).comment=Comment('Chọn mã đơn vị vận chuyển (sheet "MA DVVC"). Mã phải trùng bảng đổi mã trong app: Cài đặt → Chung → Mã ĐVVC.','HAK')
for r in range(V0,V1+1):
    v.cell(r,1,f'=IF(B{r}="","",COUNTA($B${V0}:B{r}))')
    v.cell(r,6,f'=IF(E{r}="","",IFERROR(INDEX(\'MA DVVC\'!$B$4:$B$53,MATCH(E{r},\'MA DVVC\'!$A$4:$A$53,0)),"Mã chưa có trong MA DVVC"))')
    v.cell(r,7,f'=IF(B{r}="","",COUNTIF(\'BANG CAN\'!$F${R0}:$F${R1},B{r}))')
    for ci in range(1,9):
        c=v.cell(r,ci); c.border=B; c.font=font(color='0000FF' if ci in (2,3,4,5,8) else '000000')
        if ci in (2,3,4,5,8): c.fill=fill(INPUT)
    v.cell(r,1).alignment=v.cell(r,2).alignment=v.cell(r,5).alignment=v.cell(r,7).alignment=Alignment(horizontal='center')
    v.cell(r,4).number_format='@'
v.conditional_formatting.add(f'B{V0}:B{V1}',FormulaRule(formula=[f'AND(B{V0}<>"",COUNTIF($B${V0}:$B${V1},B{V0})>1)'],font=Font(name=F,color='C92A2A',bold=True),fill=fill('FDECEC')))
v.conditional_formatting.add(f'F{V0}:F{V1}',FormulaRule(formula=[f'LEFT(F{V0},3)="Mã "'],font=Font(name=F,color='C92A2A'),fill=fill('FDECEC')))
v.conditional_formatting.add(f'G{V0}:G{V1}',FormulaRule(formula=[f'G{V0}=0'],font=Font(name=F,color='B08900')))
dvc=DataValidation(type='list',formula1="='MA DVVC'!$A$4:$A$53",allow_blank=True,errorStyle='warning',showErrorMessage=True,errorTitle='Mã ĐVVC',error='Mã này chưa có trong sheet "MA DVVC". Bổ sung mã ở sheet đó (và trong app: Cài đặt → Mã ĐVVC).')
v.add_data_validation(dvc); dvc.add(f'E{V0}:E{V1}')
dvt=DataValidation(type='custom',formula1=f'COUNTIF($B${V0}:$B${V1},B{V0})=1',allow_blank=True,errorStyle='warning',showErrorMessage=True,errorTitle='Biển số trùng',error='Biển số này đã có trong danh sách — mỗi xe chỉ 1 dòng.')
v.add_data_validation(dvt); dvt.add(f'B{V0}:B{V1}')
v.freeze_panes=f'C{V0}'; v.auto_filter.ref=f'A{HD}:H{V1}'
v.print_title_rows=f'{HD}:{HD}'; v.page_setup.orientation='landscape'; v.page_setup.paperSize=9; v.sheet_properties.pageSetUpPr.fitToPage=True; v.page_setup.fitToWidth=1; v.page_setup.fitToHeight=0
v.oddFooter.center.text='Trang &P / &N'; v.sheet_properties.tabColor=NAVY

# ================= Sheet 3: MA DVVC =================
m=wb.create_sheet('MA DVVC')
m.column_dimensions['A'].width=12; m.column_dimensions['B'].width=52
m.merge_cells('A1:B1'); m['A1']='BẢNG MÃ ĐƠN VỊ VẬN CHUYỂN'; m['A1'].font=font(size=14,bold=True,color=NAVY)
m.merge_cells('A2:B2'); m['A2']='Phải trùng với app: Cài đặt → Chung → Mã ĐVVC. Thêm đơn vị mới: ghi tiếp bên dưới (tối đa 50 dòng).'; m['A2'].font=font(size=9,italic=True,color='5B6B7C')
for i,h in enumerate(['Mã ĐVVC','Tên đơn vị vận chuyển'],1):
    c=m.cell(3,i,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(NAVY); c.alignment=CENTER; c.border=B
codes=[('HAK','Công Ty TNHH Hoàng Anh Khôi'),('HAK QN','Công Ty TNHH Hoàng Anh Khôi Quảng Nam'),('ĐH','Công Ty TNHH SX TM DV Đại Hiệp'),('HHH','CÔNG TY TNHH DỊCH VỤ VÀ VẬN TẢI HÙNG HOÀNG HOA')]
for r in range(4,54):
    for ci in (1,2):
        c=m.cell(r,ci); c.border=B; c.font=font(color='0000FF'); c.fill=fill(INPUT)
    if r-4 < len(codes): m.cell(r,1,codes[r-4][0]); m.cell(r,2,codes[r-4][1])
m.cell(4,1).comment=Comment('Mã và tên lấy theo bảng đổi mã mặc định của app (Cài đặt → Chung). Sửa cho khớp nếu công ty đã đổi trong app.','HAK')
m.sheet_properties.tabColor='868E96'

# ================= Sheet 4: HUONG DAN =================
h=wb.create_sheet('HUONG DAN'); h.column_dimensions['A'].width=4; h.column_dimensions['B'].width=110
h['B1']='HƯỚNG DẪN SỬ DỤNG FILE MẪU'; h['B1'].font=font(size=14,bold=True,color=NAVY)
lines=[
 ('h','1. Nguyên tắc'),
 ('t','• Mỗi đợt xuất hàng (1 tàu / 1 lô) = 1 file. Lưu file theo tên đợt, ví dụ: BangCan_JR-SOPHIA-V2501_16-20.08.2026.xlsx'),
 ('t','• Danh sách xe & tài xế nằm ở sheet "DS XE - TAI XE" của CHÍNH file đó và chỉ áp dụng cho bảng cân trong file — không dùng cho lô, tàu khác.'),
 ('t','• Ô nền vàng, chữ xanh = ô cần nhập. Ô nền xám / chữ đen = công thức tự tính, không sửa.'),
 ('h','2. Sheet "BANG CAN" — bảng cân xuất hàng'),
 ('t','• Ghi Đơn vị xuất, Tàu / chuyến, Kho xuất ở đầu sheet.'),
 ('t','• Mỗi phiếu cân 1 dòng: Số phiếu, Ngày cân (dd/mm/yyyy), Giờ cân vào / ra (gõ 07:05 hoặc 07:05:30), Biển số xe (chọn từ danh sách xe), Cân lần 1 (xe không), Cân lần 2 (có hàng) — đơn vị kg.'),
 ('t','• Khối lượng hàng tự tính = Cân lần 2 − Cân lần 1. Phần mềm cân đã xuất sẵn khối lượng hàng thì dán đè vào cột này.'),
 ('t','• Cột "Kiểm tra DS xe" báo đỏ "Chưa có trong DS xe" nếu biển số chưa có trong danh sách xe của đợt → bổ sung trước khi đưa vào app.'),
 ('t','• Có thể dán thẳng dữ liệu xuất từ phần mềm cân vào từ dòng 8 (giữ đúng thứ tự cột). Không chèn / xóa dòng tiêu đề (dòng 7).'),
 ('h','3. Sheet "DS XE - TAI XE" — danh sách xe của đợt'),
 ('t','• Mỗi xe 1 dòng (không trùng biển số): Biển số xe, Họ và tên tài xế, Số điện thoại, Mã ĐVVC (chọn trong danh sách).'),
 ('t','• Tên đơn vị vận chuyển tự hiện theo mã (sheet "MA DVVC"). Cột "Số chuyến trong bảng cân" đếm số phiếu của xe trong đợt (0 = xe không chạy đợt này).'),
 ('h','4. Đưa vào app QUẢN LÝ XUẤT HÀNG'),
 ('t','• Nhập liệu → chọn hình thức → thêm file này làm 1 đợt. App đọc sheet "BANG CAN" làm bảng cân.'),
 ('t','• Danh sách xe: app tự đọc sheet "DS XE - TAI XE" của cùng file và chỉ áp cho đúng đợt này — điền Tài xế, Đơn vị vận chuyển vào biên bản, BKLS, PXKKVCNB.'),
 ('t','• File cân xuất từ phần mềm cân (không có sheet DS XE): ở dòng đợt bấm 🚚 → gắn file danh sách xe riêng, hoặc chọn "Dùng danh sách của đợt khác" (mặc định đợt gần nhất). Đợt chưa có danh sách → app cảnh báo, Tài xế / ĐVVC để trống.'),
 ('t','• Tải lại file mẫu này trong app: Nhập liệu → khung "Danh sách xe & tài xế" → ⬇ Tải file Excel mẫu.'),
 ('h','5. Ví dụ 1 dòng (chỉ để tham khảo, KHÔNG nhập vào bảng)'),
]
r=3
for k,t in lines:
    c=h.cell(r,2,t); c.alignment=Alignment(wrap_text=True,vertical='top')
    c.font=font(bold=True,size=11,color=TEAL) if k=='h' else font()
    if k=='h': r+=1 if r==3 else 0
    r+=1
ex_h=['Số phiếu','Ngày cân','Giờ vào','Giờ ra','Biển số xe','Cân lần 1 (kg)','Cân lần 2 (kg)','KL hàng (kg)','Kho xuất','Ghi chú']
ex_v=['0001','16/08/2026','07:05','07:32','92C-12884','13.240','31.090','17.850','Kho HAK Đà Nẵng','JR SOPHIA V.2501']
h.cell(r,2,'BANG CAN:  '+'  |  '.join(f'{a}: {b}' for a,b in zip(ex_h,ex_v))).font=font(size=9,color='5B6B7C'); r+=1
h.cell(r,2,'DS XE - TAI XE:  Biển số xe: 92C-12884  |  Họ và tên tài xế: Nguyễn Văn An  |  Số điện thoại: 0905 123 456  |  Mã ĐVVC: HHH  →  CÔNG TY TNHH DỊCH VỤ VÀ VẬN TẢI HÙNG HOÀNG HOA').font=font(size=9,color='5B6B7C')
h.sheet_properties.tabColor='F59F00'; h.sheet_view.showGridLines=False
for s in (ws,v,m): s.sheet_view.showGridLines=True
from openpyxl.workbook.defined_name import DefinedName
# Vùng in tự co theo số dòng đã nhập (không in 1.000 dòng trống)
for idx,(sh,lastc,colk,first) in enumerate([('BANG CAN','L','F',R0),('DS XE - TAI XE','H','B',V0)]):
    q="'"+sh+"'"
    dn=DefinedName('_xlnm.Print_Area',localSheetId=idx,attr_text=f"{q}!$A$1:INDEX({q}!${lastc}:${lastc},MAX({first},{first-1}+COUNTA({q}!${colk}${first}:${colk}${first+999})))")
    wb[sh].defined_names['_xlnm.Print_Area']=dn
wb.active=0
import sys
wb.save(sys.argv[1] if len(sys.argv)>1 else '/tmp/xl/Mau_BangCan_XuatHang_va_DS_Xe.xlsx'); print('saved')
