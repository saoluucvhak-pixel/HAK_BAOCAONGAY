# -*- coding: utf-8 -*-
"""
Dựng thêm 2 template từ CÁC TEMPLATE ĐÃ CÓ trong templates/ (không cần file mẫu gốc):
  XuatKhau.docx  ← XuatKhoGui.docx  : BẢNG TỔNG HỢP HÀNG XUẤT KHẨU QUA CÂN (1 bên, tổng hợp theo ngày + phụ lục phiếu cân)
  XuatBan.docx   ← TraoDoi.docx     : BIÊN BẢN GIAO NHẬN HÀNG HÓA — LÔ HÀNG XUẤT BÁN (hợp đồng mua bán, đơn giá / thành tiền)
Chạy:  python3 tools/build_templates2.py   (rồi python3 tools/embed_templates.py)
"""
import copy, os, sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(os.path.dirname(HERE), 'templates')
sys.argv = [sys.argv[0], TPL, TPL]          # build_templates.py đọc SRC/OUT từ argv
sys.path.insert(0, HERE)
import importlib.util
spec = importlib.util.spec_from_file_location('bt', os.path.join(HERE, 'build_templates.py'))
bt = importlib.util.module_from_spec(spec)
_main = bt.__name__
src = open(os.path.join(HERE, 'build_templates.py'), encoding='utf-8').read().split("if __name__ == '__main__':")[0]
exec(compile(src, 'build_templates.py', 'exec'), bt.__dict__)
sub, seq, set_par, set_cell, loop_table, sync_grid, tc_set, all_paragraphs = bt.sub, bt.seq, bt.set_par, bt.set_cell, bt.loop_table, bt.sync_grid, bt.tc_set, bt.all_paragraphs


def tc_full(tc, text, bold=None):
    """Đặt text ô (giữ định dạng), xoá các đoạn thừa trong ô."""
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]:
        tc.remove(p)
    P = Paragraph(ps[0], None)
    set_par(P, text)
    if bold is not None and P.runs:
        P.runs[0].bold = bold


def build_xuatkhau():
    d = docx.Document(os.path.join(TPL, 'XuatKhoGui.docx'))
    sub(d, 'BẢNG XÁC NHẬN HÀNG GỬI KHO XUẤT KHẨU', 'BẢNG TỔNG HỢP HÀNG XUẤT KHẨU QUA CÂN')
    sub(d, 'Căn cứ số liệu cân xe thực tế',
        'Căn cứ số liệu cân xe thực tế tại {{kho_giao}}, hàng của {{giao_name}} được vận chuyển ra {{kho_nhan}} để nạp lên tàu xuất khẩu.')
    sub(d, 'BÊN GỬI (', 'ĐƠN VỊ XUẤT KHẨU ({{giao_role}}): {{giao_name}}')
    # bỏ khối "BÊN NHẬN" (công ty tự xuất khẩu — chỉ một bên): từ đoạn trống trước "BÊN NHẬN" đến hết khối, giữ 1 đoạn trống trước bảng
    ps = [p for p in all_paragraphs(d)]
    i0 = next(i for i, p in enumerate(ps) if p.text.startswith('BÊN NHẬN'))
    kill = [ps[i0 - 1]] + ps[i0:i0 + 4]
    assert kill[1].text.startswith('BÊN NHẬN') and kill[-1].text.startswith('Người đại diện'), [k.text for k in kill]
    for p in kill:
        p._p.getparent().remove(p._p)
    t0 = d.tables[0]
    rows = [('Tàu', '{{ship}}'), ('Thời gian xuất hàng', '{{dateFrom}} đến {{dateTo}}'), ('Kho xuất (bàn cân)', '{{kho_giao_short}}'),
            ('Cảng / nơi nạp hàng', '{{kho_nhan}}'), ('Số xe (phiếu cân)', '{{trip}}'), ('Khối lượng qua cân', '{{weightScaleTon}} tấn'),
            ('Trừ: Xe trả về', '{{returnedText}}'), ('= Khối lượng thực xuất', '{{weightDelivered}} tấn')]
    assert len(t0.rows) == len(rows) + 1
    for i, (a, b) in enumerate(rows):
        set_cell(t0, i + 1, 0, a); set_cell(t0, i + 1, 1, b)
    sub(d, 'Hai bên xác nhận', 'Số liệu trên được tổng hợp từ các phiếu cân thực tế (chi tiết tại các phụ lục), làm căn cứ theo dõi, hạch toán và đối chiếu khối lượng hàng xuất khẩu khi cần thiết.')
    # chữ ký: trái = người lập bảng (để trống ký tay), phải = đại diện đơn vị xuất khẩu
    sg = d.tables[1]
    L = [p for p in sg.rows[0].cells[0].paragraphs if p.text.strip()]
    R = [p for p in sg.rows[0].cells[1].paragraphs if p.text.strip()]
    assert len(L) == 4 and len(R) == 4
    set_par(L[0], 'NGƯỜI LẬP BẢNG')
    for p in L[1:]: set_par(p, '')
    set_par(R[0], 'ĐẠI DIỆN ĐƠN VỊ XUẤT KHẨU')
    for p, k in zip(R[1:], ('name', 'rep', 'position')): set_par(p, '{{giao_%s}}' % k)
    # --- phụ lục: 1 = tổng hợp theo ngày (bảng mới), 2 = chi tiết phiếu cân (bảng có sẵn) ---
    ps = [p for p in all_paragraphs(d)]
    h = next(p for p in ps if p.text.startswith('PHỤ LỤC:')); sh = ps[ps.index(h) + 1]
    detail = d.tables[2]
    h1, sh1 = copy.deepcopy(h._p), copy.deepcopy(sh._p)
    h._p.addprevious(h1); h._p.addprevious(sh1)
    set_par(Paragraph(h1, d), 'PHỤ LỤC 1: BẢNG TỔNG HỢP KHỐI LƯỢNG THEO NGÀY')
    set_par(Paragraph(sh1, d), '{{kho_giao_short}} → {{kho_nhan_short}} — Tàu {{ship}} — {{trip}} phiếu cân, {{weight}} tấn')
    set_par(h, 'PHỤ LỤC 2: BẢNG KÊ CHI TIẾT PHIẾU CÂN')
    dt = copy.deepcopy(detail._tbl)
    grid = dt.find(qn('w:tblGrid'))
    drop = [2, 3, 8]                                   # bỏ 3 cột: giờ cân 1, giờ cân 2, kho xuất
    for i in sorted(drop, reverse=True):
        grid.remove(grid.findall(qn('w:gridCol'))[i])
        for tr in dt.findall(qn('w:tr'))[:2]:
            tr.remove(tr.findall(qn('w:tc'))[i])
    trs = dt.findall(qn('w:tr'))
    dt.remove(trs[2])                                  # hàng tổng cũ (ô gộp) → dựng lại từ hàng dữ liệu
    heads = ['STT', 'Ngày', 'Số xe', 'Khối lượng qua cân (kg)', 'Xe trả về (kg)', 'Thực xuất (kg)']
    for tc, t in zip(trs[0].findall(qn('w:tc')), heads): tc_full(tc, t)
    fields = ['stt', 'date', 'trips', 'kg', 'retKg', 'delKg']
    dtcs = trs[1].findall(qn('w:tc'))
    for i, (tc, f) in enumerate(zip(dtcs, fields)):
        tx = '{{%s}}' % f
        if i == 0: tx = '{{#days}}' + tx
        if i == len(fields) - 1: tx += '{{/days}}'
        tc_full(tc, tx)
    tot = copy.deepcopy(trs[1]); trs[1].addnext(tot)
    for tc, t in zip(tot.findall(qn('w:tc')), ['', 'TỔNG CỘNG', '{{trip}}', '{{weightKg}}', '{{returnedKg}}', '{{weightDeliveredKg}}']):
        tc_full(tc, t, bold=True)
    sh1.addnext(dt)
    T = Table(dt, d)
    sync_grid(T, [800, 2600, 2000, 3300, 3300, 3300])
    sp = OxmlElement('w:p'); dt.addnext(sp)            # đoạn trống giữa 2 phụ lục
    d.save(os.path.join(TPL, 'XuatKhau.docx'))


def build_xuatban():
    d = docx.Document(os.path.join(TPL, 'TraoDoi.docx'))
    sub(d, 'PHỤ LỤC', 'BIÊN BẢN GIAO NHẬN HÀNG HÓA', exact=True)
    sub(d, 'BẢNG CHI TIẾT CÂN HÀNG - LÔ HÀNG TRAO ĐỔI', 'BẢNG CHI TIẾT CÂN HÀNG - LÔ HÀNG XUẤT BÁN')
    sub(d, '(Kèm theo Hợp đồng trao đổi số:', '(Theo Hợp đồng mua bán số: {{contractNo}} giữa {{giao_name}} (Bên bán) và {{nhan_name}} (Bên mua))')
    sub(d, 'Đơn vị giao:', 'Bên bán: {{giao_name}} ({{giao_short}})  —  Bên mua: {{nhan_name}} ({{nhan_short}})')
    t = d.tables[0]
    # dòng đơn giá / thành tiền ngay dưới bảng (chỉ hiện khi nhập đơn giá)
    ps = [p for p in all_paragraphs(d)]
    ref = next(p for p in ps if p.text.startswith('Bàn cân xuất:'))
    np_ = copy.deepcopy(ref._p)
    t._tbl.addnext(np_)
    P = Paragraph(np_, d); P.alignment = 0
    set_par(P, '{{#hasPrice}}Đơn giá: {{unitPrice}} đồng/tấn  —  Thành tiền: {{amount}} đồng{{priceNote}}. Bằng chữ: {{amountWords}}.{{/hasPrice}}')
    sg = d.tables[1]
    for ci, t_ in ((0, 'ĐẠI DIỆN BÊN BÁN'), (1, 'ĐẠI DIỆN BÊN MUA')):
        p0 = [p for p in sg.rows[0].cells[ci].paragraphs if p.text.strip()][0]
        set_par(p0, t_)
    d.save(os.path.join(TPL, 'XuatBan.docx'))


if __name__ == '__main__':
    build_xuatkhau(); build_xuatban()
    print('OK', sorted(os.listdir(TPL)))
